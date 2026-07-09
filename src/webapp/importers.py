"""Import standardized workbook sheets into the Web job blackboard."""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from blackboard.excel_store import ExcelBlackboardStore
from blackboard.sheet_schema import get_sheet_spec
from blackboard.validators import BlackboardValidationError, validate_row
from tools.parameter_tools import now_iso


IMPORTABLE_SHEETS = (
    "parameter_checklist",
    "project_parameters",
    "wbs_tasks_final",
    "resource_plan_final",
)

SHEET_ALIASES: dict[str, tuple[str, ...]] = {
    "parameter_checklist": ("parameter_checklist", "参数检查清单", "参数清单"),
    "project_parameters": ("project_parameters", "项目参数"),
    "wbs_tasks_final": ("wbs_tasks_final", "WBS工序分解表", "WBS", "WBS分解表"),
    "resource_plan_final": ("resource_plan_final", "资源需求表", "资源计划", "资源需求与能力"),
}

COLUMN_ALIASES: dict[str, dict[str, str]] = {
    "parameter_checklist": {
        "参数编号": "parameter_id",
        "参数类别": "category",
        "参数名称": "name",
        "是否必需": "required",
        "参数值": "value",
        "值": "value",
        "单位": "unit",
        "来源": "source",
        "状态": "status",
        "备注": "note",
    },
    "project_parameters": {
        "参数编号": "parameter_id",
        "参数值": "value",
        "值": "value",
        "单位": "unit",
        "来源": "source",
        "置信度": "confidence",
        "备注": "note",
    },
    "wbs_tasks_final": {
        "任务编号": "task_id",
        "WBS编码": "wbs_code",
        "阶段": "phase",
        "区段": "section",
        "楼层或区域": "floor_or_area",
        "任务名称": "task_name",
        "工作包": "work_package",
        "工程量": "quantity",
        "单位": "unit",
        "工期": "duration_days",
        "工期天数": "duration_days",
        "前置任务": "predecessor_ids",
        "前置任务ID": "predecessor_ids",
        "关系类型": "relation_type",
        "滞后天数": "lag_days",
        "来源": "source",
        "置信度": "confidence",
        "备注": "note",
    },
    "resource_plan_final": {
        "任务编号": "task_id",
        "资源类型": "resource_type",
        "资源名称": "resource_name",
        "需求量": "demand",
        "单位": "unit",
        "能力": "capacity",
        "资源能力": "capacity",
        "时段": "period",
        "周期": "period",
        "冲突": "conflict_flag",
        "来源": "source",
        "置信度": "confidence",
        "备注": "note",
    },
}


def import_standard_tables_from_uploads(
    input_docs_dir: Path,
    blackboard_path: Path,
) -> dict[str, Any]:
    store = ExcelBlackboardStore(blackboard_path)
    store.initialize()
    summary: dict[str, Any] = {"files_checked": 0, "imported": {}, "skipped": []}
    for path in sorted(input_docs_dir.iterdir()):
        if path.suffix.lower() not in {".xlsx", ".xlsm"}:
            continue
        summary["files_checked"] += 1
        try:
            workbook = load_workbook(path, data_only=True, read_only=True)
        except Exception as exc:
            summary["skipped"].append({"file": path.name, "reason": str(exc)})
            continue
        for sheet_name in IMPORTABLE_SHEETS:
            worksheet = _find_sheet(workbook, sheet_name)
            if worksheet is None:
                continue
            rows = _read_standard_rows(worksheet, sheet_name)
            if not rows:
                continue
            existing = store.read_rows(sheet_name)
            merged = _merge_rows(sheet_name, existing, rows)
            store.replace_rows(sheet_name, merged)
            summary["imported"][sheet_name] = len(merged)
    return summary


def _find_sheet(workbook: Any, sheet_name: str) -> Any | None:
    aliases = {normalize_header(value) for value in SHEET_ALIASES[sheet_name]}
    for candidate in workbook.worksheets:
        if normalize_header(candidate.title) in aliases:
            return candidate
    return None


def _read_standard_rows(worksheet: Any, sheet_name: str) -> list[dict[str, Any]]:
    header_values = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not header_values:
        return []
    spec = get_sheet_spec(sheet_name)
    alias_map = {normalize_header(header): header for header in spec.headers}
    alias_map.update(
        {normalize_header(alias): canonical for alias, canonical in COLUMN_ALIASES.get(sheet_name, {}).items()}
    )
    headers = [alias_map.get(normalize_header(value)) for value in header_values]
    if not any(header in spec.required_headers for header in headers if header):
        return []
    rows: list[dict[str, Any]] = []
    for values in worksheet.iter_rows(min_row=2, values_only=True):
        row = {
            header: values[index]
            for index, header in enumerate(headers)
            if header and index < len(values)
        }
        if not any(value not in {None, ""} for value in row.values()):
            continue
        row = _apply_import_defaults(sheet_name, row)
        if not all(str(row.get(header) or "").strip() for header in spec.required_headers):
            continue
        try:
            validate_row(sheet_name, row)
        except BlackboardValidationError:
            continue
        rows.append(row)
    return rows


def _apply_import_defaults(sheet_name: str, row: dict[str, Any]) -> dict[str, Any]:
    row = dict(row)
    if sheet_name == "parameter_checklist":
        row.setdefault("required", "yes")
        row.setdefault("source", "uploaded_standard_table")
        row.setdefault("extraction_status", row.get("status") or "user_confirmed")
        row.setdefault("confidence", "1.00")
        row.setdefault("status", row.get("extraction_status") or "user_confirmed")
        row.setdefault("owner_agent", "web_import")
        row.setdefault("created_by", "web_import")
        row.setdefault("created_at", now_iso())
    elif sheet_name == "project_parameters":
        row.setdefault("source", "uploaded_standard_table")
        row.setdefault("confidence", "1.00")
        row.setdefault("extraction_status", "user_confirmed")
        row.setdefault("confirmed_by", "web_import")
        row.setdefault("updated_at", now_iso())
        row.setdefault("created_by", "web_import")
    elif sheet_name == "wbs_tasks_final":
        row.setdefault("relation_type", "FS")
        row.setdefault("lag_days", 0)
        row.setdefault("source", "uploaded_standard_table")
        row.setdefault("confidence", "1.00")
        row.setdefault("owner_agent", "web_import")
    elif sheet_name == "resource_plan_final":
        row.setdefault("source", "uploaded_standard_table")
        row.setdefault("confidence", "1.00")
        row.setdefault("owner_agent", "web_import")
    return row


def _merge_rows(
    sheet_name: str,
    existing: list[dict[str, Any]],
    imported: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    key_fields = {
        "parameter_checklist": ("parameter_id",),
        "project_parameters": ("parameter_id",),
        "wbs_tasks_final": ("task_id",),
        "resource_plan_final": ("task_id", "resource_type", "resource_name", "period"),
    }[sheet_name]
    merged = {_row_key(row, key_fields): dict(row) for row in existing if _row_key(row, key_fields)}
    for row in imported:
        key = _row_key(row, key_fields)
        if key:
            merged[key] = row
    return list(merged.values())


def _row_key(row: dict[str, Any], fields: tuple[str, ...]) -> str:
    return "||".join(str(row.get(field) or "").strip() for field in fields)


def normalize_header(value: Any) -> str:
    return str(value or "").strip().replace(" ", "").replace("_", "").lower()
